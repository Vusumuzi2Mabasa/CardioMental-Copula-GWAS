"""Validate the manuscript and verify/write exact public snapshot checksums."""
from pathlib import Path
import argparse,json,hashlib,re,subprocess,zipfile
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--write',action='store_true');args=ap.parse_args()
    assert json.loads((ROOT/'results/audit.json').read_text())['passed']
    source=(ROOT/'README.md').read_text(encoding='utf-8')
    assert not re.search(r'@[A-Z_]+@',source),'Unexpanded manuscript tokens'
    document=ROOT/'report/CardioMental_Copula_GWAS_Report.docx'
    d=Document(document)
    assert d.core_properties.author=='Vusumuzi Mabasa'
    assert len(d.inline_shapes)==7
    assert [len(t.rows) for t in d.tables]==[5,5,9,9,6,7]
    with zipfile.ZipFile(document) as z:
        xml=z.read('word/document.xml').decode()
    assert xml.count('<m:oMath>')>=8
    assert all(r.target_ref.startswith(('https://','http://','#')) for r in d.part.rels.values() if r.reltype.endswith('/hyperlink'))
    words=' '.join(p.text for p in d.paragraphs)
    assert 'synthetic' in words.lower() and 'simulation' in words.lower()
    assert 'Genomic SEM analyses' in words
    # Compare each table's text against the generated research manuscript.
    tables=[];current=[]
    for line in source.splitlines()+['']:
        if line.startswith('|'):
            vals=[v.strip() for v in line.strip('|').split('|')]
            if not all(v=='---' for v in vals):current.append(vals)
        elif current:tables.append(current);current=[]
    assert len(tables)==6
    norm=lambda x:' '.join(x.split())
    for expected,actual in zip(tables,d.tables):
        got=[[norm(c.text) for c in r.cells] for r in actual.rows]
        assert got==[[norm(v) for v in row] for row in expected],'Word/source table discrepancy'
    for target in re.findall(r'\]\(([^)]+)\)',source):
        if not target.startswith(('https://','http://','#')):assert (ROOT/target).exists(),target
    sw=json.loads((ROOT/'results/word_report_software.json').read_text())
    assert sw['source_sha256']==sha(ROOT/'README.md')
    visual=json.loads((ROOT/'results/report_validation.json').read_text())
    assert visual['all_passed'] and visual['word_sha256']==sha(document)
    assert visual['source_sha256']==sha(ROOT/'README.md')
    assert visual['pages']==visual['rendered_pages_inspected']
    paths=subprocess.check_output(['git','-c','core.fsmonitor=false','ls-files','-z'],cwd=ROOT).decode().split('\0')
    paths=[p for p in paths if p]
    assert not any(p.startswith('work/') or Path(p).name.startswith('~$') for p in paths)
    assert not any(Path(p).name in ['.Renviron','.Rprofile','.Rhistory','.RData'] for p in paths)
    token=re.compile(rb'(?:gh[pousr]_[A-Za-z0-9]{25,}|github_pat_[A-Za-z0-9_]{30,}|-----BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY-----)')
    for name in paths:
        blob=(ROOT/name).read_bytes()
        if name.endswith('.docx'):
            with zipfile.ZipFile(ROOT/name) as z:blob=b'\n'.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
        assert not token.search(blob),'Possible credential content in '+name
        assert (ROOT/name).stat().st_size<90_000_000,'Oversize public file: '+name
    result=dict(passed=True,author='Vusumuzi Mabasa',figures=7,tables=6,native_math=True,
        word_tables_match_manuscript=True,portable_links=True,rendered_pages=visual['pages'],
        source_sha256=sha(ROOT/'README.md'),word_sha256=sha(document),credential_scan_passed=True)
    if args.write:
        audit_name='results/publication_audit.json'
        (ROOT/audit_name).write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8',newline='\n')
        if audit_name not in paths:paths.append(audit_name)
        lines=[f'{sha(ROOT/p)}  {p}' for p in sorted(paths) if p!='checksums_public.sha256']
        (ROOT/'checksums_public.sha256').write_text('\n'.join(lines)+'\n',encoding='utf-8',newline='\n')
        print('PASS: manuscript audit; wrote',len(lines),'public checksums')
    else:
        for line in (ROOT/'checksums_public.sha256').read_text().splitlines():
            expected,name=line.split('  ',1);assert sha(ROOT/name)==expected,name
        print('PASS: manuscript, report, public content and snapshot checksums')

if __name__=='__main__':main()
