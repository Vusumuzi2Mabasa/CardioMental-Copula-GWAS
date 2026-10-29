"""Check manuscript/report consistency, visual verification and the public manifest."""
from pathlib import Path
import sys,json,re,zipfile,subprocess,argparse,hashlib
from docx import Document
ROOT=Path(__file__).resolve().parents[2];sys.path.append(str(ROOT/'v3/src'))
from prepare import sha,dump
V=ROOT/'v4'
def paths():
    b=subprocess.check_output(['git','-c','core.fsmonitor=false','ls-files','--cached','--others','--exclude-standard','-z'],cwd=ROOT)
    return sorted(set(x for x in b.decode().split('\0') if x and x!='checksums_public.sha256' and (ROOT/x).is_file()))
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--write',action='store_true');args=ap.parse_args()
    local=json.loads((ROOT/'v3/results/audit.json').read_text());extra=json.loads((V/'results/audit.json').read_text())
    assert local['passed'] and local['summary_checks_passed'] and extra['passed']
    source=(ROOT/'README.md').read_text(encoding='utf-8');report=ROOT/'report/CardioMental_Shared_Genetics_Report_v4.docx';doc=Document(report)
    assert doc.core_properties.author=='Vusumuzi Mabasa';assert doc.core_properties.title==source.splitlines()[0].removeprefix('# ')
    assert len(doc.inline_shapes)==12 and len(doc.tables)==11
    with zipfile.ZipFile(report) as z:xml=z.read('word/document.xml').decode()
    equations=len(re.findall(r'^\$\$$',source,re.M))//2
    assert xml.count('<m:oMath>')==equations,(xml.count('<m:oMath>'),equations)
    words=' '.join(p.text for p in doc.paragraphs)
    assert 'simulated' in words.lower() and '48,000' in words and 'Genomic SEM' in words
    assert not re.search(r'@[A-Z_]+@',source)
    assert not any(token in words for token in ['\\frac','\\mathrm','\\mathrel','\\begin','\\Sigma','\\quad'])
    tables=[];current=[]
    for line in source.splitlines()+['']:
        if line.startswith('|'):
            values=[v.strip() for v in line.strip('|').split('|')]
            if not all(v=='---' for v in values):current.append(values)
        elif current:tables.append(current);current=[]
    assert len(tables)==11
    norm=lambda s:' '.join(s.split())
    for a,b in zip(tables,doc.tables):assert [[norm(c) for c in row] for row in a]==[[norm(c.text) for c in row.cells] for row in b.rows]
    for p in re.findall(r'\]\(([^)]+)\)',source):
        if not p.startswith(('https://','http://','#')):assert (ROOT/p).exists(),p
    assert all(r.target_ref.startswith(('https://','http://','#')) for r in doc.part.rels.values() if r.reltype.endswith('/hyperlink'))
    visual=json.loads((V/'results/report_validation.json').read_text())
    assert visual['passed'] and visual['pages']==visual['pages_inspected']
    assert visual['word_sha256']==sha(report) and visual['readme_sha256']==sha(ROOT/'README.md')
    assert visual['figures_inspected']==12
    provenance=json.loads((V/'results/manuscript_sources.json').read_text());assert provenance['readme_sha256']==sha(ROOT/'README.md')
    assert provenance['extension_audit_sha256']==sha(V/'results/audit.json') and provenance['local_audit_sha256']==sha(ROOT/'v3/results/audit.json')
    software=json.loads((V/'results/word_report_software.json').read_text());assert software['source_sha256']==sha(ROOT/'README.md')
    # Sensitive strings are checked in text and decompressed Word XML, not binary noise.
    credential=re.compile(rb'(?:gh[pousr]_[A-Za-z0-9]{25,}|github_pat_[A-Za-z0-9_]{30,}|-----BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY-----)')
    for p in paths():
        q=ROOT/p;assert q.stat().st_size<90*1024*1024,p
        assert '/work/' not in '/'+p and q.name not in ['.Renviron','.Rprofile','.Rhistory','.RData']
        if q.suffix.lower() in ['.py','.r','.md','.txt','.json','.yml','.csv','.toml']:
            assert not credential.search(q.read_bytes()),p
        if q.suffix.lower()=='.docx':
            with zipfile.ZipFile(q) as z:
                for name in z.namelist():
                    if name.endswith('.xml'):assert not credential.search(z.read(name)),p
    result=dict(passed=True,local_cohorts=31200,initial_polygenic_cohorts=4000,fresh_polygenic_cohorts=12800,total_new_unique_cohorts=48000,figures=12,tables=11,native_editable_equations=equations,report_pages=visual['pages'],author='Vusumuzi Mabasa',readme_sha256=sha(ROOT/'README.md'),word_sha256=sha(report),original_scientific_files_preserved=local['preserved_old_files'],credential_scan_passed=True)
    dump(V/'results/publication_audit.json',result)
    public=paths();manifest=ROOT/'checksums_public.sha256'
    expected=''.join(sha(ROOT/p)+'  '+p+'\n' for p in public)
    if args.write:manifest.write_text(expected,encoding='utf-8',newline='\n')
    assert manifest.read_text(encoding='utf-8')==expected
    print(json.dumps(dict(**result,public_files=len(public)),indent=2))
if __name__=='__main__':main()
