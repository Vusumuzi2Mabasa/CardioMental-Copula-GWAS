"""Check shared manuscript, baseline preservation and complete public manifest."""
from pathlib import Path
import json,hashlib,re,zipfile,subprocess,argparse
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--write',action='store_true');args=ap.parse_args()
    audit=json.loads((ROOT/'v2/results/audit.json').read_text());assert audit['passed'] and audit['baseline_bytes_preserved']
    text=(ROOT/'README.md').read_text(encoding='utf-8');assert not re.search(r'@[A-Z_]+@',text)
    document=ROOT/'report/CardioMental_Copula_GWAS_Report_v2.docx';d=Document(document)
    assert d.core_properties.author=='Vusumuzi Mabasa'
    assert len(d.inline_shapes)==6 and [len(t.rows) for t in d.tables]==[5,9,9,6]
    with zipfile.ZipFile(document) as z:xml=z.read('word/document.xml').decode()
    assert xml.count('<m:oMath>')>=10
    words=' '.join(p.text for p in d.paragraphs)
    assert 'synthetic' in words.lower() and 'Genomic SEM analyses' in words and '92,000' in words
    assert all(r.target_ref.startswith(('https://','http://','#')) for r in d.part.rels.values() if r.reltype.endswith('/hyperlink'))
    tables=[];current=[]
    for line in text.splitlines()+['']:
        if line.startswith('|'):
            vals=[v.strip() for v in line.strip('|').split('|')]
            if not all(v=='---' for v in vals):current.append(vals)
        elif current:tables.append(current);current=[]
    assert len(tables)==4
    norm=lambda s:' '.join(s.split())
    for source,actual in zip(tables,d.tables):assert [[norm(v) for v in row] for row in source]==[[norm(c.text) for c in row.cells] for row in actual.rows]
    for link in re.findall(r'\]\(([^)]+)\)',text):
        if not link.startswith(('http://','https://','#')):assert (ROOT/link).exists(),link
    visual=json.loads((ROOT/'v2/results/report_validation.json').read_text())
    assert visual['all_passed'] and visual['pages']==visual['rendered_pages_inspected']
    assert visual['source_sha256']==sha(ROOT/'README.md') and visual['word_sha256']==sha(document)
    sw=json.loads((ROOT/'v2/results/word_report_software.json').read_text());assert sw['source_sha256']==sha(ROOT/'README.md')
    paths=subprocess.check_output(['git','-c','core.fsmonitor=false','ls-files','-z'],cwd=ROOT).decode().split('\0');paths=[p for p in paths if p]
    token=re.compile(rb'(?:gh[pousr]_[A-Za-z0-9]{25,}|github_pat_[A-Za-z0-9_]{30,}|-----BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY-----)')
    def scan(blob):assert not token.search(blob),'Possible credential in public snapshot'
    for path in paths:
        p=Path(path)
        assert not any(part in ['work','__pycache__','.venv'] for part in p.parts)
        assert p.name not in ['.Renviron','.Rprofile','.Rhistory','.RData'] and not p.name.startswith('~$')
        assert (ROOT/path).stat().st_size<90_000_000,path
        if path.endswith(('.docx','.zip')):
            with zipfile.ZipFile(ROOT/path) as z:
                for n in z.namelist():
                    if not n.endswith(('.png','.docx')):scan(z.read(n))
        else:scan((ROOT/path).read_bytes())
    result=dict(passed=True,source_sha256=sha(ROOT/'README.md'),word_sha256=sha(document),word_tables_match_README=True,
                figures=6,tables=4,native_math=True,author='Vusumuzi Mabasa',rendered_pages=visual['pages'],
                baseline_preserved=True,credential_scan_passed=True)
    if args.write:
        name='v2/results/publication_audit.json';(ROOT/name).write_text(json.dumps(result,indent=2)+'\n')
        if name not in paths:paths.append(name)
        (ROOT/'checksums_public.sha256').write_text('\n'.join(f'{sha(ROOT/p)}  {p}' for p in sorted(paths) if p!='checksums_public.sha256')+'\n')
        print('PASS wrote publication audit and',len(paths)-1,'public checksums')
    else:
        names=[]
        for line in (ROOT/'checksums_public.sha256').read_text().splitlines():
            expected,name=line.split('  ',1);assert sha(ROOT/name)==expected,name;names.append(name)
        assert set(names)==set(paths)-{'checksums_public.sha256'},'Public manifest missing tracked files'
        print('PASS manuscript Word baseline credential scan and all public snapshot hashes')
if __name__=='__main__':main()
